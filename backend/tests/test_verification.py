from app.verification import evaluate_liveness


def test_liveness_requires_a_real_frame_sequence():
    result = evaluate_liveness([])
    assert result["status"] == "INCONCLUSIVE"
    assert result["head_turn"] is False
    assert result["blink"] is False
    assert "four" in result["reason"]


def test_biometric_key_is_required_for_encryption(monkeypatch):
    monkeypatch.delenv("ATTENDANCE_BIOMETRIC_ENCRYPTION_KEY", raising=False)
    from app.verification import VerificationUnavailable, encrypt_embedding

    try:
        encrypt_embedding([0.1, 0.2])
    except VerificationUnavailable as error:
        assert "ENCRYPTION_KEY" in str(error)
    else:
        raise AssertionError("Biometric encryption must fail closed without a server key")


def test_workbook_contains_auditable_sheets():
    from openpyxl import load_workbook
    from app.reports import build_attendance_workbook

    workbook = load_workbook(build_attendance_workbook({"session": {"id": "s1"}, "summary": {}, "rows": []}, []))
    assert workbook.sheetnames == [
        "Executive Summary",
        "Session Attendance",
        "Student-wise Attendance",
        "Verification Review",
        "Corrections and Audit History",
    ]
