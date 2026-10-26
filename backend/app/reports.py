"""Workbook generation for authorized attendance exports."""
from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _write_table(sheet: Any, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B554F")
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(48, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width


def build_attendance_workbook(report: dict[str, Any], audit_items: list[dict[str, Any]]) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Executive Summary"
    session = report.get("session") or {}
    counts = report.get("summary") or {}
    summary.append(["Attendrix attendance export"])
    summary["A1"].font = Font(bold=True, size=16, color="0B554F")
    summary.append(["Generated at", report.get("generated_at", "")])
    summary.append(["Session", session.get("title", "")])
    summary.append(["Class", session.get("class_name", "")])
    summary.append(["Session ID", session.get("id", "")])
    summary.append([])
    _write_table(summary, ["Metric", "Value"], ((key.replace("_", " ").title(), value) for key, value in counts.items()))

    rows = report.get("rows") or []
    attendance = workbook.create_sheet("Session Attendance")
    _write_table(
        attendance,
        ["Student ID", "Student Name", "Email", "Submission Time", "Verification Status", "Review Decision", "Final Attendance Status", "Note", "Updated At"],
        ((row.get("student_code", ""), row.get("name", ""), row.get("email", ""), row.get("submission_time", ""), row.get("verification_status", ""), row.get("review_decision", ""), row.get("status", ""), row.get("note", ""), row.get("updated_at", "")) for row in rows),
    )

    student = workbook.create_sheet("Student-wise Attendance")
    _write_table(student, ["Student ID", "Student Name", "Status", "Verification Status", "Updated At"], ((row.get("student_code", ""), row.get("name", ""), row.get("status", ""), row.get("verification_status", ""), row.get("updated_at", "")) for row in rows))

    review = workbook.create_sheet("Verification Review")
    _write_table(review, ["Student ID", "Student Name", "Submission Time", "Verification Status", "Review Decision", "Reason"], ((row.get("student_code", ""), row.get("name", ""), row.get("submission_time", ""), row.get("verification_status", ""), row.get("review_decision", ""), row.get("note", "")) for row in rows))

    audit = workbook.create_sheet("Corrections and Audit History")
    _write_table(audit, ["Timestamp", "Action", "Actor", "Reason", "Metadata"], ((item.get("timestamp", ""), item.get("action", ""), item.get("actor_name", item.get("actor_id", "")), item.get("note", ""), item.get("metadata_json", "")) for item in audit_items))

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
